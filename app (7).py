# app.py
import io
import os
import re
import openpyxl
import pandas as pd
import requests
import streamlit as st
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

# Регистрация шрифтов с поддержкой кириллицы для PDF
font_registered = False
for f_path in [
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "DejaVuSans.ttf"
]:
    if os.path.exists(f_path):
        try:
            pdfmetrics.registerFont(TTFont("DejaVuSans", f_path))
            b_path = f_path.replace("DejaVuSans.ttf", "DejaVuSans-Bold.ttf")
            if os.path.exists(b_path):
                pdfmetrics.registerFont(TTFont("DejaVuSans-Bold", b_path))
            else:
                pdfmetrics.registerFont(TTFont("DejaVuSans-Bold", f_path))
            font_registered = True
            break
        except Exception:
            pass

# ----------------------------------------------------------------------------
# 1. НАСТРОЙКА СТРАНИЦЫ
# ----------------------------------------------------------------------------
st.set_page_config(
    page_title="Брусника — сравнение участников тендера",
    page_icon="🌲",
    layout="wide",
    initial_sidebar_state="collapsed"
)

# ----------------------------------------------------------------------------
# 2. API КЛЮЧ
# ----------------------------------------------------------------------------
api_key = st.secrets.get("FOCUS_API_KEY", "")
if not api_key:
    st.warning("⚠️ API ключ Контур.Фокуса не настроен. Добавьте FOCUS_API_KEY в secrets приложения.")

# ЦВЕТА
BG_COLOR = "#f8f7f4"
WHITE = "#ffffff"
DARK = "#1e1e1e"
GRAY = "#6b6b6b"
LIGHT_GRAY = "#f0eeea"
BORDER_COLOR = "rgba(0, 0, 0, 0.05)"
GOLD = "#b89b7b"
GREEN = "#1d6b2e"
RED = "#b33a3a"
ORANGE = "#b8681a"
BLUE = "#1e5a8a"

# CSS
st.markdown(f"""
<style>
    .reportview-container {{
        background: {BG_COLOR};
    }}
    .main-header {{
        padding: 1.5rem 0 0.5rem 0;
        border-bottom: 1px solid {BORDER_COLOR};
        margin-bottom: 2rem;
    }}
    .brand-tag {{
        font-size: 0.85rem;
        letter-spacing: 0.15em;
        text-transform: uppercase;
        color: {GOLD};
        font-weight: 600;
        margin-bottom: 0.25rem;
    }}
    .tender-card {{
        background: {WHITE};
        border-radius: 12px;
        padding: 1.5rem;
        border: 1px solid {BORDER_COLOR};
        margin-bottom: 1.5rem;
        box-shadow: 0 2px 8px rgba(0,0,0,0.02);
    }}
    .winner-card {{
        border-left: 5px solid {GREEN};
        background: linear-gradient(90deg, rgba(29, 107, 46, 0.04) 0%, {WHITE} 100%);
    }}
    .reserve-card {{
        border-left: 5px solid {ORANGE};
        background: linear-gradient(90deg, rgba(184, 104, 26, 0.04) 0%, {WHITE} 100%);
    }}
    .other-card {{
        border-left: 5px solid {GRAY};
    }}
    .badge {{
        display: inline-block;
        padding: 0.25rem 0.6rem;
        border-radius: 4px;
        font-size: 0.8rem;
        font-weight: 600;
    }}
    .badge-winner {{ background-color: rgba(29, 107, 46, 0.15); color: {GREEN}; }}
    .badge-reserve {{ background-color: rgba(184, 104, 26, 0.15); color: {ORANGE}; }}
    .badge-other {{ background-color: {LIGHT_GRAY}; color: {DARK}; }}
</style>
""", unsafe_allow_html=True)

# ----------------------------------------------------------------------------
# 3. ФУНКЦИЯ РАЗБОРА EXCEL-ТАБЛИЦЫ
# ----------------------------------------------------------------------------
def parse_tender_excel(file_bytes_or_path):
    """
    Разбирает сводную таблицу коммерческих предложений (как ТУ, так и ТТ).
    Возвращает dict:
      - 'tender_title': название тендера/предмет торгов
      - 'lots': список названий лотов
      - 'lots_count': количество лотов
      - 'participants': список dict [{'inn': str, 'name': str, 'prices': [float, ...]}]
    """
    if isinstance(file_bytes_or_path, str):
        wb = openpyxl.load_workbook(file_bytes_or_path, data_only=True)
    elif isinstance(file_bytes_or_path, bytes):
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes_or_path), data_only=True)
    else:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes_or_path.getvalue()), data_only=True)
    
    ws = wb.active

    # Извлечение темы тендера
    tender_title = ""
    for r in range(1, 10):
        for c in range(1, 10):
            val = ws.cell(r, c).value
            if val and any(k in str(val).lower() for k in ["предмет тендера", "предмет торгов", "конкурентный лист"]):
                neighbor = ws.cell(r, c + 1).value or ws.cell(r, c + 2).value or val
                tender_title = str(neighbor).strip()
                break
        if tender_title:
            break

    # 1. Поиск участников и их ИНН
    inn_row = None
    vendor_row = None
    for r in range(1, min(15, ws.max_row + 1)):
        for c in range(1, min(15, ws.max_column + 1)):
            v = ws.cell(r, c).value
            if v:
                s_val = str(v).strip().upper()
                if s_val == "ИНН":
                    inn_row = r
                    break
                elif "НАИМЕНОВАНИЕ КОНТРАГЕНТА" in s_val or "НАИМЕНОВАНИЕ ПОСТАВЩИКА" in s_val:
                    vendor_row = r
        if inn_row:
            break

    participants = []
    if inn_row:
        # Формат ТУ: отдельная строка с "ИНН"
        for c in range(1, ws.max_column + 1):
            val = ws.cell(inn_row, c).value
            if val is not None:
                digits = re.findall(r"\b\d{10}\b|\b\d{12}\b", str(val).strip())
                if digits:
                    inn = digits[0]
                    name_val = ws.cell(inn_row - 1, c).value or f"Участник {len(participants)+1}"
                    participants.append({
                        'col': c,
                        'inn': inn,
                        'name': str(name_val).strip()
                    })
    else:
        # Формат ТТ: строка с наименованием поставщиков и ИНН в скобках
        target_row = vendor_row if vendor_row else 7
        for c in range(1, ws.max_column + 1):
            val = ws.cell(target_row, c).value
            if val:
                s_val = str(val).strip()
                match = re.search(r"ИНН\s*[:\s]?\s*(\d{10}|\d{12})", s_val, re.IGNORECASE)
                if match:
                    inn = match.group(1)
                    name_clean = re.sub(r"\(?\s*ИНН.*?\)??", "", s_val).strip()
                    participants.append({
                        'col': c,
                        'inn': inn,
                        'name': name_clean or f"Участник {len(participants)+1}"
                    })

    participants.sort(key=lambda x: x['col'])

    # 2. Поиск строк лотов
    lots = []
    for r in range(1, ws.max_row + 1):
        for c in (1, 2, 3, 4):
            val = ws.cell(r, c).value
            if val:
                s_val = str(val).strip()
                if s_val.lower().startswith("стоимость на"):
                    continue
                match = re.search(r"Лот\s*(?:№|N)?\s*(\d+)", s_val, re.IGNORECASE)
                if match:
                    lot_num = int(match.group(1))
                    if not any(l['lot_num'] == lot_num for l in lots):
                        lots.append({
                            'row': r,
                            'lot_num': lot_num,
                            'lot_name': s_val
                        })
                    break

    lots.sort(key=lambda x: x['lot_num'])

    if not lots:
        lots = [{'row': 21, 'lot_num': 1, 'lot_name': 'Лот №1'}]

    # 3. Извлечение цен по каждому лоту
    result_participants = []
    for p in participants:
        c = p['col']
        prices = []
        for l in lots:
            r = l['row']
            price = 0.0
            for cand in [c + 1, c, c + 2]:
                val = ws.cell(r, cand).value
                if isinstance(val, (int, float)):
                    price = float(val)
                    break
                elif val is not None:
                    cleaned = str(val).replace(" ", "").replace(",", ".").replace("\xa0", "")
                    try:
                        price = float(cleaned)
                        break
                    except ValueError:
                        pass
            prices.append(price)

        result_participants.append({
            'inn': p['inn'],
            'name': p['name'],
            'prices': prices
        })

    return {
        'tender_title': tender_title or "Сводная таблица коммерческих предложений",
        'lots': [l['lot_name'] for l in lots],
        'lots_count': len(lots),
        'participants': result_participants
    }

# ----------------------------------------------------------------------------
# 4. ШАПКА
# ----------------------------------------------------------------------------
st.markdown("""
<div class="main-header">
    <div class="brand-tag">🌲 Брусника · закупки и тендеры</div>
    <h1 style="margin: 0; font-size: 2.1rem; color: #1e1e1e;">Сравнение и рейтингование участников закупки</h1>
    <p style="color: #6b6b6b; margin-top: 0.5rem; font-size: 1.05rem;">
        Комплексная оценка надежности и ценовых предложений по данным ФНС, бухотчетности и Контур.Фокуса
    </p>
</div>
""", unsafe_allow_html=True)

# ----------------------------------------------------------------------------
# 5. ИНИЦИАЛИЗАЦИЯ СОСТОЯНИЯ
# ----------------------------------------------------------------------------
if 'lots_count' not in st.session_state:
    st.session_state.lots_count = 1
if 'lots_names' not in st.session_state:
    st.session_state.lots_names = ["Лот №1"]
if 'tender_title' not in st.session_state:
    st.session_state.tender_title = "Тендер"
if 'participants' not in st.session_state:
    st.session_state.participants = [
        {"inn": "", "name": "", "prices": [0.0] * st.session_state.lots_count},
        {"inn": "", "name": "", "prices": [0.0] * st.session_state.lots_count},
    ]
if 'analysis_result' not in st.session_state:
    st.session_state.analysis_result = None
if 'run_analysis' not in st.session_state:
    st.session_state.run_analysis = False

# ----------------------------------------------------------------------------
# 6. БЛОК ЗАГРУЗКИ EXCEL И АВТОЗАПОЛНЕНИЯ
# ----------------------------------------------------------------------------
st.markdown("### 📥 1. Загрузка сводной таблицы Excel")
uploaded_file = st.file_uploader(
    "Загрузите сводную оценочную таблицу или конкурентный лист (.xlsx)",
    type=["xlsx"],
    help="Файл будет автоматически разобран: определятся все лоты, участники, их ИНН и суммы предложений."
)

if uploaded_file is not None:
    col_up1, col_up2 = st.columns([1, 4])
    with col_up1:
        parse_btn = st.button("⚡ Разобрать таблицу", type="secondary", use_container_width=True)
    
    if parse_btn or 'last_uploaded_name' not in st.session_state or st.session_state.last_uploaded_name != uploaded_file.name:
        try:
            parsed = parse_tender_excel(uploaded_file)
            st.session_state.last_uploaded_name = uploaded_file.name
            st.session_state.tender_title = parsed['tender_title']
            st.session_state.lots_count = parsed['lots_count']
            st.session_state.lots_names = parsed['lots']
            st.session_state.participants = parsed['participants']
            st.session_state.analysis_result = None
            st.success(
                f"✅ Успешно разобрано: **{parsed['lots_count']} лот(ов)** и "
                f"**{len(parsed['participants'])} участников**! Значения автоматически заполнены."
            )
            st.rerun()
        except Exception as e:
            st.error(f"Ошибка при разборе файла Excel: {str(e)}")

# ----------------------------------------------------------------------------
# 7. ФУНКЦИИ РАБОТЫ С API КОНТУР.ФОКУСА
# ----------------------------------------------------------------------------
def fetch_kontur_data(inn, api_key, custom_name=""):
    """Собирает данные из всех необходимых методов Контур.Фокуса."""
    result = {
        'inn': inn,
        'name': custom_name or 'Неизвестно',
        'revenue': 0,
        'employees': 0,
        'yellow_statements': 0,
        'red_statements': 0,
        'risk_text': [],
        'risk_markers': [],
        'brief_href': '',
        'scoring_score': 0.0,
        'max_debt': 0,
        'cred_day': 0,
        'equity': 0,
        'status': 'OK'
    }

    if not api_key:
        result['status'] = 'No API Key'
        return result

    SCORING_MODEL_IDS = {
        "304a90a4-8882-4be8-9c01-72e2b356179d",
        "7f5311e1-e7c0-4535-a85b-4bb2eebcee79",
        "3e64aa23-2653-412d-b072-c04a948e6402",
    }

    try:
        # 1. /api3/req — Название
        r = requests.get("https://focus-api.kontur.ru/api3/req", params={"key": api_key, "inn": inn}, timeout=15)
        if r.status_code == 200 and r.json():
            req = r.json()[0]
            fetched_name = req.get('UL', {}).get('legalName', {}).get('short') or req.get('IP', {}).get('fio')
            if fetched_name:
                result['name'] = fetched_name

        # 2. /api3/accountingReports — Выручка (код 2110)
        r = requests.get("https://focus-api.kontur.ru/api3/accountingReports", params={"key": api_key, "inn": inn}, timeout=15)
        buh_forms = []
        if r.status_code == 200 and r.json():
            data = r.json()
            if data and len(data) > 0:
                buh_forms = data[0].get('buhForms', [])
                years = sorted([f.get('year') for f in buh_forms if f.get('year')], reverse=True)
                if years:
                    latest = next((f for f in buh_forms if f.get('year') == years[0]), None)
                    if latest:
                        for item in latest.get('form2', []):
                            if item.get('code') == 2110:
                                result['revenue'] = item.get('endValue', 0)
                                break

        # 3. /api3/legalAnalytics — Штат
        try:
            r = requests.get("https://focus-api.kontur.ru/api3/legalAnalytics", params={"key": api_key, "inn": inn}, timeout=15)
            if r.status_code == 200 and r.json():
                data = r.json()
                entry = data[0] if isinstance(data, list) and data else data
                legal = entry.get('legalAnalyticsData', {}) if isinstance(entry, dict) else {}
                staff_block = legal.get('staffNumber')
                if isinstance(staff_block, dict):
                    val = staff_block.get('data')
                    if val is not None:
                        try:
                            result['employees'] = int(val)
                        except (TypeError, ValueError):
                            pass
                if result['employees'] == 0:
                    hist_block = legal.get('historyStaffNumbers')
                    if isinstance(hist_block, dict):
                        hist = hist_block.get('data', []) or []
                        if hist:
                            def _d(rec):
                                return rec.get('date', '') if isinstance(rec, dict) else ''
                            last = sorted(hist, key=_d)[-1]
                            v = last.get('staffNumber') if isinstance(last, dict) else None
                            if v is not None:
                                try:
                                    result['employees'] = int(v)
                                except (TypeError, ValueError):
                                    pass
        except Exception:
            pass

        # 4. /api3/briefReport — флаги + ссылка
        r = requests.get("https://focus-api.kontur.ru/api3/briefReport", params={"key": api_key, "inn": inn}, timeout=15)
        if r.status_code == 200 and r.json():
            data = r.json()
            brief = data[0] if isinstance(data, list) else data
            brief_report = brief.get('briefReport', {}) or {}
            summary = brief_report.get('summary', {}) or {}
            href = brief_report.get('href', '')
            has_yellow = bool(summary.get('yellowStatements', False))
            has_red = bool(summary.get('redStatements', False))
            result['yellow_statements'] = 1 if has_yellow else 0
            result['red_statements'] = 1 if has_red else 0
            result['brief_href'] = href

        # 5. /api3/scoring — сумма + Risk-маркеры
        try:
            r = requests.get("https://focus-api.kontur.ru/api3/scoring", params={"key": api_key, "inn": inn}, timeout=15)
            if r.status_code == 200 and r.json():
                data = r.json()
                scoring_list = data if isinstance(data, list) else [data]
                scoring_sum = 0
                risk_markers = []
                for entry in scoring_list:
                    if not isinstance(entry, dict):
                        continue
                    models = (
                        entry.get('scoringData', []) or entry.get('scoring', []) or entry.get('models', [])
                    )
                    for model in models:
                        if not isinstance(model, dict):
                            continue
                        model_id = model.get('modelId') or model.get('id')
                        if model_id in SCORING_MODEL_IDS:
                            rating = model.get('rating') or model.get('score') or 0
                            try:
                                scoring_sum += float(rating)
                            except (TypeError, ValueError):
                                pass
                        for m in model.get('triggeredMarkers', []) or []:
                            if not isinstance(m, dict):
                                continue
                            if m.get('impact') != 'Risk':
                                continue
                            name = m.get('name', '')
                            desc = m.get('description', '') or ''
                            weight = m.get('weight', '')
                            txt = name + (f" — {desc}" if desc else "")
                            if weight in ('High', 'Significant'):
                                risk_markers.append(('red', txt))
                            else:
                                risk_markers.append(('yellow', txt))
                result['scoring_score'] = round(scoring_sum, 2)
                result['risk_markers'] = risk_markers
        except Exception:
            pass

        for kind, txt in result.get('risk_markers', []):
            if kind == 'red':
                result['risk_text'].append(f"🔴 {txt}")
            else:
                result['risk_text'].append(f"⚠️ {txt}")

        if result['yellow_statements'] and not any(t.startswith('⚠️') for t in result['risk_text']):
            result['risk_text'].append("⚠️ Есть информация, помеченная жёлтым цветом")
        if result['red_statements'] and not any(t.startswith('🔴') for t in result['risk_text']):
            result['risk_text'].append("🔴 Есть информация, помеченная красным цветом")

        if result.get('brief_href'):
            result['risk_text'].append(
                f'<a href="{result["brief_href"]}" target="_blank">🔗 Открыть полный отчёт Контур.Фокуса</a>'
            )

        # 6-8. Бухотчетность
        if buh_forms:
            years = sorted([f.get('year') for f in buh_forms if f.get('year')], reverse=True)
            last_year = years[0] if years else None

            def get_val(form, code, vtype='endValue'):
                for item in form:
                    if item.get('code') == code:
                        return item.get(vtype, 0)
                return 0

            if last_year:
                latest = next((f for f in buh_forms if f.get('year') == last_year), None)
                if latest:
                    f1 = latest.get('form1', [])
                    f2 = latest.get('form2', [])
                    eV_1230 = get_val(f1, 1230)
                    eV_2110 = get_val(f2, 2110)
                    eV_1520 = get_val(f1, 1520)
                    eV_1600 = get_val(f1, 1600)
                    eV_1400 = get_val(f1, 1400)
                    eV_1500 = get_val(f1, 1500)

                    NDS = 22
                    K1 = 1 + NDS / 100
                    K2 = 0.8
                    result['max_debt'] = max(0, int(((eV_1230 + eV_2110 * K1 * K2) - 0) / 12 * 0.6))
                    if eV_2110 > 0:
                        result['cred_day'] = int(((eV_1520 + eV_1520) / 2) / eV_2110 * 365)
                    result['equity'] = eV_1600 - (eV_1400 + eV_1500)

        return result
    except Exception as e:
        result['status'] = f'Error: {str(e)}'
        return result

# ----------------------------------------------------------------------------
# 8. МЕТОДИКА РАСЧЕТА БАЛЛОВ ПО ЛОТУ
# ----------------------------------------------------------------------------
def calculate_lot_scores(df_all, lot_idx, lots_count):
    """
    Считает баллы для одного лота.
    Участники с ценой по лоту > 0.
    """
    df = df_all.copy()
    df['lot_price'] = df['lots_prices'].apply(
        lambda prices: prices[lot_idx] if lot_idx < len(prices) else 0.0
    )
    df = df[df['lot_price'] > 0].reset_index(drop=True)
    if df.empty:
        return df

    df['score'] = 0
    df['score_details'] = [[] for _ in range(len(df))]

    # 1. Наименьшая доля тендера к выручке
    df['revenue_ratio'] = df.apply(
        lambda r: (r['price_total'] / r['revenue'] * 100) if r['revenue'] > 0 else 999,
        axis=1
    )
    s = df['revenue_ratio'].sort_values().index.tolist()
    if len(s) >= 1:
        df.loc[s[0], 'score'] += 2
        df.loc[s[0], 'score_details'].append("Доля тендера к выручке: +2")
    if len(s) >= 2:
        df.loc[s[1], 'score'] += 1
        df.loc[s[1], 'score_details'].append("Доля тендера к выручке: +1")

    # 2. Наибольший аванс
    s = df['max_debt'].sort_values(ascending=False).index.tolist()
    if len(s) >= 1:
        df.loc[s[0], 'score'] += 2
        df.loc[s[0], 'score_details'].append("Наибольший аванс: +2")
    if len(s) >= 2:
        df.loc[s[1], 'score'] += 1
        df.loc[s[1], 'score_details'].append("Наибольший аванс: +1")

    # 3. Превышение аванса к 10%
    df['advance_excess'] = df.apply(
        lambda r: (r['max_debt'] / (r['price_total'] * 0.1) * 100) if r['price_total'] > 0 else 0,
        axis=1
    )
    s = df['advance_excess'].sort_values(ascending=False).index.tolist()
    if len(s) >= 1:
        df.loc[s[0], 'score'] += 2
        df.loc[s[0], 'score_details'].append("Превышение аванса к 10%: +2")
    if len(s) >= 2:
        df.loc[s[1], 'score'] += 1
        df.loc[s[1], 'score_details'].append("Превышение аванса к 10%: +1")

    # 4. Нет рисковых записей
    for idx, row in df.iterrows():
        if row['yellow_statements'] == 0 and row['red_statements'] == 0:
            df.loc[idx, 'score'] += 2
            df.loc[idx, 'score_details'].append("Нет рисковых записей: +2")

    # 5. Наибольший штат
    s = df['employees'].sort_values(ascending=False).index.tolist()
    if len(s) >= 1:
        df.loc[s[0], 'score'] += 2
        df.loc[s[0], 'score_details'].append("Наибольший штат: +2")
    if len(s) >= 2:
        df.loc[s[1], 'score'] += 1
        df.loc[s[1], 'score_details'].append("Наибольший штат: +1")

    # 6. Наибольший скоринг
    s = df['scoring_score'].sort_values(ascending=False).index.tolist()
    if len(s) >= 1:
        df.loc[s[0], 'score'] += 2
        df.loc[s[0], 'score_details'].append("Наилучший скоринг: +2")
    if len(s) >= 2:
        df.loc[s[1], 'score'] += 1
        df.loc[s[1], 'score_details'].append("Наилучший скоринг: +1")

    # 7. Наименьшая оборачиваемость
    df['cred_day_pos'] = df['cred_day'].apply(lambda x: x if x > 0 else 9999)
    s = df['cred_day_pos'].sort_values().index.tolist()
    valid = [i for i in s if df.loc[i, 'cred_day_pos'] < 9999]
    if len(valid) >= 1:
        df.loc[valid[0], 'score'] += 2
        df.loc[valid[0], 'score_details'].append("Наименьшая оборачиваемость: +2")
    if len(valid) >= 2:
        df.loc[valid[1], 'score'] += 1
        df.loc[valid[1], 'score_details'].append("Наименьшая оборачиваемость: +1")

    # 8. Наибольшие чистые активы
    s = df['equity'].sort_values(ascending=False).index.tolist()
    if len(s) >= 1:
        df.loc[s[0], 'score'] += 2
        df.loc[s[0], 'score_details'].append("Наибольшие чистые активы: +2")
    if len(s) >= 2:
        df.loc[s[1], 'score'] += 1
        df.loc[s[1], 'score_details'].append("Наибольшие чистые активы: +1")

    # 9. Наименьшая цена по КОНКРЕТНОМУ ЛОТУ
    s = df['lot_price'].sort_values().index.tolist()
    if len(s) >= 1:
        df.loc[s[0], 'score'] += 6
        df.loc[s[0], 'score_details'].append("Наименьшая цена по лоту: +6")
    if len(s) >= 2:
        df.loc[s[1], 'score'] += 3
        df.loc[s[1], 'score_details'].append("Наименьшая цена по лоту: +3")

    df = df.sort_values('score', ascending=False).reset_index(drop=True)
    return df

# ----------------------------------------------------------------------------
# 9. ГЕНЕРАЦИЯ ОБЩЕГО PDF-ОТЧЕТА
# ----------------------------------------------------------------------------
def build_pdf_report(tender_summary, lots_results):
    """
    Формирует PDF-отчет со сводными результатами и лотами.
    """
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=20,
        rightMargin=20,
        topMargin=20,
        bottomMargin=20
    )
    
    styles = getSampleStyleSheet()
    font_fam = "DejaVuSans" if font_registered else "Helvetica"
    font_bold_fam = "DejaVuSans-Bold" if font_registered else "Helvetica-Bold"

    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName=font_bold_fam,
        fontSize=15,
        leading=18,
        textColor=colors.HexColor('#1e1e1e')
    )
    h2_style = ParagraphStyle(
        'H2',
        parent=styles['Normal'],
        fontName=font_bold_fam,
        fontSize=11,
        leading=15,
        textColor=colors.HexColor('#1e5a8a'),
        spaceBefore=8,
        spaceAfter=5
    )
    normal_style = ParagraphStyle(
        'NormalRu',
        parent=styles['Normal'],
        fontName=font_fam,
        fontSize=8,
        leading=11,
        textColor=colors.HexColor('#222222')
    )
    table_header_style = ParagraphStyle(
        'TH',
        parent=styles['Normal'],
        fontName=font_bold_fam,
        fontSize=7,
        leading=9,
        textColor=colors.white,
        alignment=1
    )
    table_cell_style = ParagraphStyle(
        'TC',
        parent=styles['Normal'],
        fontName=font_fam,
        fontSize=7,
        leading=9,
        textColor=colors.HexColor('#222222')
    )
    badge_green = ParagraphStyle(
        'BG',
        parent=styles['Normal'],
        fontName=font_bold_fam,
        fontSize=7,
        leading=9,
        textColor=colors.HexColor('#1d6b2e')
    )
    badge_red = ParagraphStyle(
        'BR',
        parent=styles['Normal'],
        fontName=font_fam,
        fontSize=7,
        leading=9,
        textColor=colors.HexColor('#b33a3a')
    )
    
    story = []
    
    # Заголовок
    story.append(Paragraph("<b>Брусника</b> · Сводный отчет по результатам оценки тендера", title_style))
    story.append(Spacer(1, 4))
    t_name = tender_summary.get('tender_name', 'Закупка')
    story.append(Paragraph(f"<b>Предмет тендера:</b> {t_name}", normal_style))
    story.append(Paragraph(
        f"<b>Количество лотов:</b> {tender_summary.get('lots_count', 0)} | "
        f"<b>Количество участников:</b> {tender_summary.get('participants_count', 0)}",
        normal_style
    ))
    story.append(Spacer(1, 8))
    
    # 1. Итоговый рейтинг
    story.append(Paragraph("1. Итоговый рейтинг участников тендера", h2_style))
    overall_df = tender_summary.get('overall_ranking')
    if overall_df is not None and not overall_df.empty:
        table_data = [[
            Paragraph("Место", table_header_style),
            Paragraph("Участник", table_header_style),
            Paragraph("ИНН", table_header_style),
            Paragraph("Сумма по всем лотам, ₽", table_header_style),
            Paragraph("Побед в лотах", table_header_style),
            Paragraph("Сумма баллов", table_header_style),
            Paragraph("Выручка, ₽", table_header_style),
            Paragraph("Штат, чел", table_header_style),
            Paragraph("Скоринг", table_header_style),
            Paragraph("Аванс, ₽", table_header_style),
            Paragraph("Риски (К/Ж)", table_header_style)
        ]]
        
        for idx, row in overall_df.iterrows():
            pos = f"#{idx+1}"
            if idx == 0:
                pos = "🏆 1"
            elif idx == 1:
                pos = "🥈 2"
            elif idx == 2:
                pos = "🥉 3"
                
            risk_str = f"{row.get('red_statements', 0)} / {row.get('yellow_statements', 0)}"
            risk_style = badge_red if row.get('red_statements', 0) > 0 else normal_style
            
            table_data.append([
                Paragraph(pos, table_cell_style),
                Paragraph(str(row.get('name', ''))[:35], table_cell_style),
                Paragraph(str(row.get('inn', '')), table_cell_style),
                Paragraph(f"{row.get('price_total', 0):,.0f}".replace(",", " "), table_cell_style),
                Paragraph(str(row.get('won_lots', 0)), table_cell_style),
                Paragraph(f"<b>{row.get('total_score', 0)}</b>", badge_green),
                Paragraph(f"{row.get('revenue', 0):,.0f}".replace(",", " "), table_cell_style),
                Paragraph(str(row.get('employees', 0)), table_cell_style),
                Paragraph(str(row.get('scoring_score', 0)), table_cell_style),
                Paragraph(f"{row.get('max_debt', 0):,.0f}".replace(",", " "), table_cell_style),
                Paragraph(risk_str, risk_style)
            ])
            
        t = Table(table_data, colWidths=[40, 160, 65, 95, 60, 65, 80, 50, 50, 80, 55], repeatRows=1)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e5a8a')),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#dddddd')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8f7f4')]),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
            ('TOPPADDING', (0, 0), (-1, -1), 3),
        ]))
        story.append(t)
        
    story.append(Spacer(1, 10))
    story.append(PageBreak())
    story.append(Paragraph("2. Результаты оценки предложений по лотам", h2_style))
    story.append(Spacer(1, 4))
    
    for lot in lots_results:
        lot_idx = lot['lot_idx']
        lot_name = lot.get('lot_name', f"Лот №{lot_idx + 1}")
        df_lot = lot.get('df_lot')
        
        lot_story = []
        lot_story.append(Paragraph(f"<b>{lot_name}</b>", h2_style))
        
        if df_lot is None or df_lot.empty:
            lot_story.append(Paragraph("<i>На данный лот не подано заявок (все цены = 0).</i>", normal_style))
        else:
            lot_table_data = [[
                Paragraph("Место", table_header_style),
                Paragraph("Участник", table_header_style),
                Paragraph("ИНН", table_header_style),
                Paragraph("Цена лота, ₽", table_header_style),
                Paragraph("Баллы", table_header_style),
                Paragraph("Детализация баллов", table_header_style)
            ]]
            
            for idx, row in df_lot.iterrows():
                pos = f"#{idx+1}"
                if idx == 0:
                    pos = "🏆 Победитель"
                elif idx == 1:
                    pos = "🔄 Резерв"
                    
                score_details = row.get('score_details', [])
                details_text = "; ".join(score_details) if score_details else "—"
                
                lot_table_data.append([
                    Paragraph(pos, table_cell_style),
                    Paragraph(str(row.get('name', ''))[:35], table_cell_style),
                    Paragraph(str(row.get('inn', '')), table_cell_style),
                    Paragraph(f"{row.get('lot_price', 0):,.0f}".replace(",", " "), table_cell_style),
                    Paragraph(f"<b>{row.get('score', 0)}</b>", badge_green if idx == 0 else table_cell_style),
                    Paragraph(details_text, table_cell_style)
                ])
                
            lot_t = Table(lot_table_data, colWidths=[65, 170, 70, 95, 45, 355], repeatRows=1)
            lot_t.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2c3e50')),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#dddddd')),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8f7f4')]),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
                ('TOPPADDING', (0, 0), (-1, -1), 3),
            ]))
            lot_story.append(lot_t)
            lot_story.append(Spacer(1, 8))
            
        story.append(KeepTogether(lot_story))
        
    doc.build(story)
    pdf_bytes = buffer.getvalue()
    buffer.close()
    return pdf_bytes

def render_card(row, position_label, css_class, badge_class, lot_idx, lots_count):
    """Карточка участника в контексте конкретного лота."""
    risk_html = ""
    if row['risk_text']:
        risks_joined = '<br>'.join(row['risk_text'][:10])
        risk_html = (
            f'<div style="margin-top: 0.75rem; padding: 0.5rem; background: #fff5f5; border-radius: 6px; font-size: 0.85rem;">'
            f'<strong>Рисковые записи:</strong><br>{risks_joined}'
            f'</div>'
        )

    lot_price = row['lots_prices'][lot_idx] if lot_idx < len(row['lots_prices']) else 0
    details_text = ", ".join(row.get('score_details', [])) or "Базовые условия"

    html = (
        f'<div class="tender-card {css_class}">'
        f'<div style="display: flex; justify-content: space-between; align-items: flex-start; margin-bottom: 0.75rem;">'
        f'<div>'
        f'<h4 style="margin: 0; color: #1e1e1e;">{position_label} {row["name"]}</h4>'
        f'<span style="color: #6b6b6b; font-size: 0.85rem;">ИНН {row["inn"]}</span>'
        f'</div>'
        f'<div class="badge badge-{badge_class}">'
        f'Баллы: <strong>{row["score"]}</strong>'
        f'</div>'
        f'</div>'
        f'<div style="display: grid; grid-template-columns: repeat(auto-fit, minmax(180px, 1fr)); gap: 0.75rem; font-size: 0.9rem;">'
        f'<div><strong>Цена по лоту:</strong><br><span style="font-size: 1.1rem; color: #1e5a8a; font-weight: 600;">{lot_price:,.0f} ₽</span></div>'
        f'<div><strong>Сумма по всем лотам:</strong><br>{row["price_total"]:,.0f} ₽</div>'
        f'<div><strong>Выручка (посл. год):</strong><br>{row["revenue"]:,.0f} ₽</div>'
        f'<div><strong>Допустимый аванс:</strong><br>{row["max_debt"]:,.0f} ₽</div>'
        f'<div><strong>Оборачиваемость:</strong><br>{row["cred_day"]} дн.</div>'
        f'<div><strong>Чистые активы:</strong><br>{row["equity"]:,.0f} ₽</div>'
        f'<div><strong>Скоринг / Штат:</strong><br>{row["scoring_score"]} / {row["employees"]} чел.</div>'
        f'</div>'
        f'<div style="margin-top: 0.75rem; font-size: 0.85rem; color: #555;">'
        f'<em>Обоснование баллов:</em> {details_text}'
        f'</div>'
        f'{risk_html}'
        f'</div>'
    )
    st.markdown(html, unsafe_allow_html=True)

# ----------------------------------------------------------------------------
# 10. ФОРМА: КОЛИЧЕСТВО ЛОТОВ
# ----------------------------------------------------------------------------
st.markdown("### 🧩 2. Параметры закупки")
col_l1, col_l2 = st.columns([1, 2])
with col_l1:
    lots_count = st.number_input(
        "Количество лотов в закупке",
        min_value=1,
        max_value=50,
        value=int(st.session_state.lots_count),
        step=1,
        key="lots_count_input",
        help="Количество лотов. При загрузке Excel вычисляется автоматически."
    )

if int(lots_count) != st.session_state.lots_count:
    old_n = st.session_state.lots_count
    new_n = int(lots_count)
    for p in st.session_state.participants:
        old_prices = p.get("prices", [0.0] * old_n)
        if new_n > old_n:
            p["prices"] = old_prices + [0.0] * (new_n - old_n)
        else:
            p["prices"] = old_prices[:new_n]
    if len(st.session_state.lots_names) < new_n:
        st.session_state.lots_names += [f"Лот №{k+1}" for k in range(len(st.session_state.lots_names), new_n)]
    else:
        st.session_state.lots_names = st.session_state.lots_names[:new_n]
    st.session_state.lots_count = new_n
    st.session_state.analysis_result = None
    st.rerun()

# ----------------------------------------------------------------------------
# 11. ФОРМА: УЧАСТНИКИ И ЦЕНЫ
# ----------------------------------------------------------------------------
st.markdown("### 👥 3. Участники тендера и цены по лотам")

for i, participant in enumerate(st.session_state.participants):
    with st.expander(f"Участник {i + 1}: {participant.get('name') or 'Новый контрагент'} (ИНН: {participant.get('inn') or '—'})", expanded=True):
        col_name, col_inn, col_del = st.columns([3, 2, 1])
        with col_name:
            p_name = st.text_input(
                "Наименование организации",
                value=participant.get("name", ""),
                key=f"name_{i}",
                placeholder='ООО "Компания"'
            )
            st.session_state.participants[i]["name"] = p_name
        with col_inn:
            inn_val = st.text_input(
                "ИНН",
                value=participant["inn"],
                max_chars=12,
                placeholder="10 или 12 цифр",
                key=f"inn_{i}"
            )
            st.session_state.participants[i]["inn"] = inn_val
        with col_del:
            st.write("")
            if len(st.session_state.participants) > 1:
                if st.button("🗑️ Удалить", key=f"del_{i}", use_container_width=True):
                    st.session_state.participants.pop(i)
                    st.session_state.analysis_result = None
                    st.rerun()

        st.markdown("**Цены по лотам (₽):**")
        lot_cols_per_row = 3
        prices = participant.get("prices", [0.0] * st.session_state.lots_count)
        if len(prices) != st.session_state.lots_count:
            prices = (prices + [0.0] * st.session_state.lots_count)[:st.session_state.lots_count]
            st.session_state.participants[i]["prices"] = prices

        for row_start in range(0, st.session_state.lots_count, lot_cols_per_row):
            cols = st.columns(lot_cols_per_row)
            for j in range(lot_cols_per_row):
                lot_idx = row_start + j
                if lot_idx >= st.session_state.lots_count:
                    break
                with cols[j]:
                    l_title = st.session_state.lots_names[lot_idx] if lot_idx < len(st.session_state.lots_names) else f"Лот {lot_idx + 1}"
                    short_l_title = l_title[:35] + ("..." if len(l_title) > 35 else "")
                    val = st.number_input(
                        short_l_title,
                        value=float(prices[lot_idx]),
                        min_value=0.0,
                        step=10000.0,
                        format="%.2f",
                        key=f"price_{i}_{lot_idx}"
                    )
                    st.session_state.participants[i]["prices"][lot_idx] = val

        total = sum(st.session_state.participants[i]["prices"])
        st.markdown(f"**Итоговая сумма по заявке:** `{total:,.2f} ₽`")

col_add, col_check = st.columns([1, 2])
with col_add:
    if st.button("➕ Добавить участника", use_container_width=True, key="add_participant"):
        st.session_state.participants.append({
            "name": "",
            "inn": "",
            "prices": [0.0] * st.session_state.lots_count
        })
        st.session_state.analysis_result = None
        st.rerun()

with col_check:
    if st.button("🚀 Запустить оценку и рейтингование", use_container_width=True, type="primary", key="submit_check"):
        errors = []
        for idx, p in enumerate(st.session_state.participants):
            if not p["inn"] or len(p["inn"]) not in [10, 12] or not p["inn"].isdigit():
                errors.append(f"Участник {idx + 1}: некорректный ИНН ('{p['inn']}')")
            if sum(p["prices"]) <= 0:
                errors.append(f"Участник {idx + 1}: не указано ни одной цены по лотам")
        if errors:
            for err in errors:
                st.error(f"❌ {err}")
        else:
            st.session_state.run_analysis = True
            st.rerun()

# ----------------------------------------------------------------------------
# 12. ВЫПОЛНЕНИЕ АНАЛИЗА
# ----------------------------------------------------------------------------
if st.session_state.run_analysis:
    st.session_state.run_analysis = False
    with st.spinner("Запрос данных из Контур.Фокуса и проверка участников..."):
        rows = []
        for p in st.session_state.participants:
            data = fetch_kontur_data(p['inn'], api_key, custom_name=p.get('name', ''))
            data['lots_prices'] = list(p['prices'])
            data['price_total'] = sum(p['prices'])
            rows.append(data)
        df = pd.DataFrame(rows)
        st.session_state.analysis_result = df

# ----------------------------------------------------------------------------
# 13. ВЫВОД ИТОГОВОГО СРАВНЕНИЯ, РЕЙТИНГА И РЕЗУЛЬТАТОВ ПО ЛОТАМ
# ----------------------------------------------------------------------------
if st.session_state.analysis_result is not None:
    df_all = st.session_state.analysis_result
    lots_count = st.session_state.lots_count
    
    st.markdown("---")
    st.markdown("## 🏆 Итоговое сравнение и рейтингование участников тендера")
    
    lots_results = []
    participant_scores_map = {row['inn']: 0 for _, row in df_all.iterrows()}
    participant_wins_map = {row['inn']: 0 for _, row in df_all.iterrows()}

    for lot_idx in range(lots_count):
        lot_name = st.session_state.lots_names[lot_idx] if lot_idx < len(st.session_state.lots_names) else f"Лот №{lot_idx + 1}"
        df_lot = calculate_lot_scores(df_all, lot_idx, lots_count)
        lots_results.append({
            'lot_idx': lot_idx,
            'lot_name': lot_name,
            'df_lot': df_lot
        })
        if not df_lot.empty:
            winner_inn = df_lot.iloc[0]['inn']
            participant_wins_map[winner_inn] = participant_wins_map.get(winner_inn, 0) + 1
            for _, r in df_lot.iterrows():
                participant_scores_map[r['inn']] = participant_scores_map.get(r['inn'], 0) + r['score']

    # Формируем сводный рейтинг
    overall_df = df_all.copy()
    overall_df['total_score'] = overall_df['inn'].map(participant_scores_map)
    overall_df['won_lots'] = overall_df['inn'].map(participant_wins_map)
    overall_df = overall_df.sort_values(
        by=['won_lots', 'total_score', 'price_total'],
        ascending=[False, False, True]
    ).reset_index(drop=True)

    # Топ-3
    col_rank_1, col_rank_2, col_rank_3 = st.columns(3)
    if len(overall_df) > 0:
        with col_rank_1:
            top1 = overall_df.iloc[0]
            st.success(
                f"🥇 **1 место**: {top1['name']}\n\n"
                f"• Побед в лотах: **{top1['won_lots']}** из {lots_count}\n\n"
                f"• Суммарно баллов: **{top1['total_score']}**\n\n"
                f"• Общая цена: **{top1['price_total']:,.0f} ₽**"
            )
    if len(overall_df) > 1:
        with col_rank_2:
            top2 = overall_df.iloc[1]
            st.info(
                f"🥈 **2 место**: {top2['name']}\n\n"
                f"• Побед в лотах: **{top2['won_lots']}** из {lots_count}\n\n"
                f"• Суммарно баллов: **{top2['total_score']}**\n\n"
                f"• Общая цена: **{top2['price_total']:,.0f} ₽**"
            )
    if len(overall_df) > 2:
        with col_rank_3:
            top3 = overall_df.iloc[2]
            st.warning(
                f"🥉 **3 место**: {top3['name']}\n\n"
                f"• Побед в лотах: **{top3['won_lots']}** из {lots_count}\n\n"
                f"• Суммарно баллов: **{top3['total_score']}**\n\n"
                f"• Общая цена: **{top3['price_total']:,.0f} ₽**"
            )

    # Сводная таблица
    st.markdown("#### Сводная аналитическая таблица тендера")
    summary_cols = ['name', 'inn', 'won_lots', 'total_score', 'price_total', 'revenue', 'employees', 'scoring_score', 'max_debt', 'cred_day', 'equity', 'red_statements', 'yellow_statements']
    display_overall = overall_df[summary_cols].copy()
    display_overall.columns = [
        'Участник', 'ИНН', 'Побед в лотах', 'Сумма баллов', 'Сумма предложений',
        'Выручка', 'Штат', 'Скоринг', 'Аванс', 'Оборач. (дн)', 'Чистые активы', 'Красных рисков', 'Жёлтых рисков'
    ]
    st.dataframe(
        display_overall.style.format({
            'Сумма предложений': '{:,.0f} ₽',
            'Выручка': '{:,.0f} ₽',
            'Аванс': '{:,.0f} ₽',
            'Чистые активы': '{:,.0f} ₽'
        }),
        use_container_width=True,
        hide_index=True
    )

    # Экспорт PDF
    st.markdown("### 📄 Экспорт отчета по тендеру")
    tender_summary = {
        'tender_name': st.session_state.tender_title,
        'lots_count': lots_count,
        'participants_count': len(overall_df),
        'overall_ranking': overall_df
    }

    try:
        pdf_bytes = build_pdf_report(tender_summary, lots_results)
        st.download_button(
            label="📥 Скачать общий отчет по тендеру (PDF)",
            data=pdf_bytes,
            file_name="tender_evaluation_report.pdf",
            mime="application/pdf",
            type="primary"
        )
    except Exception as e:
        st.error(f"Не удалось сформировать PDF-отчет: {str(e)}")

    # ----------------------------------------------------------------------------
    # 14. ДЕТАЛЬНЫЕ РЕЗУЛЬТАТЫ ПО КАЖДОМУ ЛОТУ
    # ----------------------------------------------------------------------------
    st.markdown("---")
    st.markdown("## 📊 Детальные результаты по лотам")
    
    for lot_data in lots_results:
        lot_idx = lot_data['lot_idx']
        lot_name = lot_data['lot_name']
        df_lot = lot_data['df_lot']

        st.markdown(f"### 🧩 {lot_name}")
        if df_lot.empty:
            st.info(f"На {lot_name} не подано ни одной заявки (все цены = 0).")
            continue

        display_cols = ['name', 'inn', 'lot_price', 'price_total', 'revenue', 'employees', 'scoring_score', 'max_debt', 'cred_day', 'equity', 'yellow_statements', 'red_statements', 'score']
        display_df = df_lot[display_cols].copy()
        display_df.columns = [
            'Компания', 'ИНН', 'Цена по лоту', 'Сумма по всем лотам', 'Выручка',
            'Штат', 'Скоринг (сумма)', 'Аванс', 'Оборач.', 'Чистые активы', 'Жёлтых', 'Красных', 'Баллы'
        ]
        st.dataframe(
            display_df.style.format({
                'Цена по лоту': '{:,.0f} ₽',
                'Сумма по всем лотам': '{:,.0f} ₽',
                'Выручка': '{:,.0f} ₽',
                'Аванс': '{:,.0f} ₽',
                'Чистые активы': '{:,.0f} ₽'
            }),
            use_container_width=True,
            hide_index=True
        )

        winner = df_lot.iloc[0]
        render_card(winner, "🏆 Победитель лота: ", "winner-card", "winner", lot_idx, lots_count)

        if len(df_lot) > 1:
            reserve = df_lot.iloc[1]
            render_card(reserve, "🔄 Резерв: ", "reserve-card", "reserve", lot_idx, lots_count)

        if len(df_lot) > 2:
            with st.expander(f"Остальные предложения по лоту ({len(df_lot) - 2})"):
                for idx in range(2, len(df_lot)):
                    row = df_lot.iloc[idx]
                    render_card(row, f"#{idx + 1}", "other-card", "other", lot_idx, lots_count)

# ----------------------------------------------------------------------------
# 15. ПОДВАЛ
# ----------------------------------------------------------------------------
st.markdown("""
<div style="margin-top: 4rem; padding: 2rem 0; text-align: center; color: #888; font-size: 0.85rem; border-top: 1px solid rgba(0,0,0,0.05);">
    © 2026 Брусника · Автоматизированная оценка надежности контрагентов и сравнение предложений
</div>
""", unsafe_allow_html=True)
